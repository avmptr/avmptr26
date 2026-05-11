import io
import json
import os
from datetime import datetime
from pathlib import Path
from typing import List, Optional

import psycopg2
import psycopg2.extras
from fastapi import FastAPI, HTTPException, Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

DATABASE_URL = os.environ.get("DATABASE_URL", "")

CONDITION_CODES = ["V1", "V2", "V3", "G1", "G2", "G3", "S", "Stg", "K", "R", "M", "Sb"]


def get_db() -> psycopg2.extensions.connection:
    url = DATABASE_URL
    if url.startswith("postgres://"):
        url = url.replace("postgres://", "postgresql://", 1)
    return psycopg2.connect(url, cursor_factory=psycopg2.extras.RealDictCursor)


def row_to_dict(row: dict) -> dict:
    d = dict(row)
    if "created_at" in d and hasattr(d["created_at"], "isoformat"):
        d["created_at"] = d["created_at"].isoformat()
    if "conditions" in d and isinstance(d["conditions"], str):
        d["conditions"] = json.loads(d["conditions"])
    return d


def init_db():
    conn = get_db()
    cur = conn.cursor()

    cur.execute("""
        CREATE TABLE IF NOT EXISTS trees (
            id          SERIAL PRIMARY KEY,
            tree_code   TEXT NOT NULL UNIQUE,
            block       TEXT NOT NULL,
            row_num     INTEGER NOT NULL,
            plot_num    INTEGER NOT NULL,
            variety     TEXT NOT NULL
        )
    """)
    cur.execute("""
        CREATE TABLE IF NOT EXISTS observations (
            id          SERIAL PRIMARY KEY,
            tree_code   TEXT NOT NULL,
            block       TEXT NOT NULL,
            conditions  TEXT NOT NULL,
            notes       TEXT,
            observed_at TEXT NOT NULL,
            created_at  TIMESTAMP NOT NULL DEFAULT NOW(),
            FOREIGN KEY (tree_code) REFERENCES trees(tree_code)
        )
    """)
    cur.execute("CREATE INDEX IF NOT EXISTS idx_obs_tree    ON observations(tree_code)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_obs_date    ON observations(observed_at)")
    cur.execute("CREATE INDEX IF NOT EXISTS idx_trees_block ON trees(block)")
    conn.commit()

    # Auto-seed trees on first deploy
    cur.execute("SELECT COUNT(*) AS n FROM trees")
    if cur.fetchone()["n"] == 0:
        from seed import TREES
        psycopg2.extras.execute_batch(
            cur,
            "INSERT INTO trees (tree_code, block, row_num, plot_num, variety) "
            "VALUES (%s, %s, %s, %s, %s) ON CONFLICT DO NOTHING",
            TREES,
        )
        conn.commit()
        print(f"Auto-seeded {len(TREES)} trees")

    cur.close()
    conn.close()


app = FastAPI(title="Meisari Avocado Farm", version="1.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

init_db()


# --- Schemas ---

class ObservationIn(BaseModel):
    tree_code: str
    conditions: List[str]
    notes: Optional[str] = None
    observed_at: Optional[str] = None  # YYYY-MM-DD


# --- Routes ---

@app.get("/blocks")
def list_blocks():
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT DISTINCT block FROM trees ORDER BY block")
    rows = cur.fetchall()
    cur.close(); conn.close()
    return [r["block"] for r in rows]


@app.get("/trees")
def list_trees(block: Optional[str] = Query(None)):
    conn = get_db()
    cur = conn.cursor()
    if block:
        cur.execute(
            "SELECT tree_code, block, row_num, plot_num, variety "
            "FROM trees WHERE block = %s ORDER BY row_num, plot_num",
            (block,),
        )
    else:
        cur.execute(
            "SELECT tree_code, block, row_num, plot_num, variety "
            "FROM trees ORDER BY block, row_num, plot_num"
        )
    rows = cur.fetchall()
    cur.close(); conn.close()
    return [dict(r) for r in rows]


@app.get("/condition-codes")
def list_condition_codes():
    return CONDITION_CODES


@app.post("/observations", status_code=201)
def create_observation(obs: ObservationIn):
    if not obs.conditions:
        raise HTTPException(400, "At least one condition code is required")

    invalid = [c for c in obs.conditions if c not in CONDITION_CODES]
    if invalid:
        raise HTTPException(400, f"Invalid condition codes: {invalid}")

    conn = get_db()
    cur = conn.cursor()

    cur.execute("SELECT block FROM trees WHERE tree_code = %s", (obs.tree_code,))
    tree = cur.fetchone()
    if not tree:
        cur.close(); conn.close()
        raise HTTPException(404, f"Tree '{obs.tree_code}' not found")

    observed_at = obs.observed_at or datetime.now().date().isoformat()

    cur.execute(
        "INSERT INTO observations (tree_code, block, conditions, notes, observed_at) "
        "VALUES (%s, %s, %s, %s, %s) RETURNING *",
        (obs.tree_code, tree["block"], json.dumps(obs.conditions), obs.notes, observed_at),
    )
    row = cur.fetchone()
    conn.commit()
    cur.close(); conn.close()
    return row_to_dict(row)


@app.get("/observations")
def list_observations(
    block: Optional[str] = Query(None),
    tree_code: Optional[str] = Query(None),
    date_from: Optional[str] = Query(None),
    date_to: Optional[str] = Query(None),
    limit: int = Query(100, le=500),
):
    conn = get_db()
    cur = conn.cursor()

    query = "SELECT * FROM observations WHERE 1=1"
    params: list = []

    if block:
        query += " AND block = %s"
        params.append(block)
    if tree_code:
        query += " AND tree_code = %s"
        params.append(tree_code)
    if date_from:
        query += " AND observed_at >= %s"
        params.append(date_from)
    if date_to:
        query += " AND observed_at <= %s"
        params.append(date_to)

    query += " ORDER BY created_at DESC LIMIT %s"
    params.append(limit)

    cur.execute(query, params)
    rows = cur.fetchall()
    cur.close(); conn.close()
    return [row_to_dict(r) for r in rows]


@app.get("/stats")
def stats():
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT block, COUNT(*) AS tree_count FROM trees GROUP BY block ORDER BY block")
    tree_rows = cur.fetchall()
    cur.execute("SELECT block, COUNT(*) AS obs_count FROM observations GROUP BY block ORDER BY block")
    obs_rows = cur.fetchall()
    cur.close(); conn.close()

    obs_map = {r["block"]: r["obs_count"] for r in obs_rows}
    return [
        {"block": r["block"], "trees": r["tree_count"], "observations": obs_map.get(r["block"], 0)}
        for r in tree_rows
    ]


@app.get("/export", summary="Download observations as Excel")
def export_observations(
    block: Optional[str] = Query(None),
    date_from: Optional[str] = Query(None, description="YYYY-MM-DD"),
    date_to: Optional[str] = Query(None, description="YYYY-MM-DD"),
):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    conn = get_db()
    cur = conn.cursor()

    query = """
        SELECT
            o.observed_at,
            o.tree_code,
            o.block,
            t.variety,
            o.conditions,
            o.notes,
            o.created_at
        FROM observations o
        JOIN trees t ON t.tree_code = o.tree_code
        WHERE 1=1
    """
    params: list = []
    if block:
        query += " AND o.block = %s"
        params.append(block)
    if date_from:
        query += " AND o.observed_at >= %s"
        params.append(date_from)
    if date_to:
        query += " AND o.observed_at <= %s"
        params.append(date_to)
    query += " ORDER BY o.observed_at DESC, o.created_at DESC"

    cur.execute(query, params)
    rows = cur.fetchall()
    cur.close()

    # Summary stats per block
    cur2 = conn.cursor()
    cur2.execute("""
        SELECT
            o.block,
            COUNT(*) AS total_obs,
            COUNT(DISTINCT o.tree_code) AS trees_observed
        FROM observations o
        WHERE 1=1
    """ + (" AND o.block = %s" if block else "") +
    (" AND o.observed_at >= %s" if date_from else "") +
    (" AND o.observed_at <= %s" if date_to else "") +
    " GROUP BY o.block ORDER BY o.block",
    [p for p in [block, date_from, date_to] if p])
    summary_rows = cur2.fetchall()
    cur2.close()
    conn.close()

    wb = Workbook()

    # ── Sheet 1: Observations ──────────────────────────────────────
    ws = wb.active
    ws.title = "Observations"

    header_fill = PatternFill("solid", fgColor="2D6A2D")
    header_font = Font(bold=True, color="FFFFFF")
    headers = ["Date", "Tree Code", "Block", "Variety", "Conditions", "Notes", "Submitted At"]
    col_widths = [12, 14, 8, 10, 30, 30, 20]

    for col, (h, w) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[cell.column_letter].width = w

    for row in rows:
        conditions = json.loads(row["conditions"]) if isinstance(row["conditions"], str) else row["conditions"]
        created_at = row["created_at"].isoformat() if hasattr(row["created_at"], "isoformat") else str(row["created_at"])
        ws.append([
            row["observed_at"],
            row["tree_code"],
            row["block"],
            row["variety"],
            ", ".join(conditions),
            row["notes"] or "",
            created_at,
        ])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{len(rows) + 1}"

    # ── Sheet 2: Summary ───────────────────────────────────────────
    ws2 = wb.create_sheet("Summary")
    ws2.column_dimensions["A"].width = 10
    ws2.column_dimensions["B"].width = 18
    ws2.column_dimensions["C"].width = 18

    sum_headers = ["Block", "Total Observations", "Trees Observed"]
    for col, h in enumerate(sum_headers, 1):
        cell = ws2.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row in summary_rows:
        ws2.append([row["block"], row["total_obs"], row["trees_observed"]])

    # Condition code breakdown
    ws2.cell(row=len(summary_rows) + 3, column=1, value="Condition Code Legend").font = Font(bold=True)
    legend = [
        ("V1", "Umur 0–1 bulan"),     ("V2", "Umur 1–12 bulan"),    ("V3", "Umur >12 bulan"),
        ("G1", "Berbunga"),            ("G2", "Buah kecil/muda"),     ("G3", "Buah dewasa/panen"),
        ("S",  "Sakit"),               ("Stg","Stagnan"),             ("K",  "Kritis"),
        ("R",  "Recovery"),            ("M",  "Mati"),                ("Sb", "Substitusi"),
    ]
    for i, (code, desc) in enumerate(legend):
        r = len(summary_rows) + 4 + i
        ws2.cell(row=r, column=1, value=code).font = Font(bold=True)
        ws2.cell(row=r, column=2, value=desc)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    date_str = datetime.now().strftime("%Y%m%d")
    filename = f"meisari_farm_{date_str}.xlsx"
    if block:
        filename = f"meisari_farm_blok{block}_{date_str}.xlsx"

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@app.get("/health")
def health():
    return {"status": "ok"}


if Path("static").exists():
    app.mount("/", StaticFiles(directory="static", html=True), name="static")
